"""Excel file the admin hands to a union / transporter to fill in.

Made so that almost nothing has to be typed:
  * mobile number: the only must; must be 10 digits (Excel refuses anything else)
  * name: Hindi or English, both fine
  * city: pick from a list (all districts when the pincode directory is loaded). Not there? type the
    city / town name, or fill the 6-digit pincode instead. Nothing is refused.
  * vehicles: one column per vehicle, pick "हाँ" from a list. No commas, no spelling.
  * owners: number of vehicles must be a whole number
The second sheet explains it in Hindi and English with an example. Only the first sheet is imported.
"""
import io

from openpyxl import Workbook
from openpyxl.comments import Comment
from openpyxl.styles import Alignment, Border, Font, PatternFill, Side
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.datavalidation import DataValidation

from .places import _data as _places

ROWS = 1000
YES = "हाँ"
# order = how they appear in the file
VEHICLE_COLS = [("truck", "ट्रक"), ("trailer", "ट्रेलर"), ("bus", "बस"), ("pickup", "पिकअप"),
                ("jcb", "जेसीबी"), ("tractor", "ट्रैक्टर"), ("car", "कार"), ("auto", "ऑटो")]

TEAL = PatternFill("solid", fgColor="0E5A6B")
YELLOW = PatternFill("solid", fgColor="FFF1D6")
SOFT = PatternFill("solid", fgColor="DCEEF1")
WHITE_BOLD = Font(bold=True, color="FFFFFF", size=12)
THIN = Side(style="thin", color="DCE5E7")


def city_options(districts: list[tuple[str, str]] = ()) -> list[str]:
    """"रीवा (Rewa), Madhya Pradesh" — the importer understands this form.
    `districts` (district, state) from the pincode directory add every other district as "Sidhi, Madhya Pradesh"."""
    d = _places()
    opts = [f"{c['hi']} ({c['en']}), {c['state']}" for c in d["cities"]]
    known = {(c["en"].lower(), c["state"]) for c in d["cities"]}
    opts += [f"{dist}, {st}" for dist, st in districts if dist and st and (dist.lower(), st) not in known]
    return sorted(set(opts))


def build(role: str, districts: list[tuple[str, str]] = ()) -> bytes:
    wb = Workbook()
    ws = wb.active
    ws.title = "भरें - Fill here"

    cols: list[tuple[str, int, str]] = [("मोबाइल नंबर (10 अंक) *", 22, "phone"), ("नाम", 24, "name")]
    if role == "owner":
        cols.append(("फर्म / ट्रांसपोर्ट का नाम", 26, "business"))
    cols.append(("शहर (सूची से चुनें या लिखें)", 34, "city"))
    cols.append(("पिनकोड (शहर न मिले तो)", 16, "pincode"))
    if role == "owner":
        cols.append(("कुल गाड़ियाँ (गिनती)", 18, "count"))
    first_vehicle = len(cols) + 1
    for _, hi in VEHICLE_COLS:
        cols.append((hi, 10, "vehicle"))

    # header
    for i, (title, width, kind) in enumerate(cols, start=1):
        c = ws.cell(row=1, column=i, value=title)
        # vehicle columns in yellow: "pick हाँ here"
        c.font = Font(bold=True, color="2A1C00", size=12) if kind == "vehicle" else WHITE_BOLD
        c.fill = PatternFill("solid", fgColor="F4A21A") if kind == "vehicle" else TEAL
        c.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        ws.column_dimensions[get_column_letter(i)].width = width
    ws.row_dimensions[1].height = 36
    ws.freeze_panes = "B2"

    # a band above the vehicle columns would shift rows; use a comment on the first vehicle header instead
    ws.cell(row=1, column=first_vehicle).comment = Comment(
        "जो गाड़ी चलाते / रखते हैं, उसके नीचे \"हाँ\" चुनें। बाकी खाली छोड़ें।\n"
        "Pick \"हाँ\" under each vehicle they drive / own. Leave the rest empty.", "Vahanza")

    last = ROWS + 1
    letter = {kind: get_column_letter(i) for i, (_, _, kind) in enumerate(cols, start=1) if kind != "vehicle"}

    # mobile: exactly 10 digits, kept as text so nothing gets rounded or turned into 9.87E+09
    phone = DataValidation(type="textLength", operator="equal", formula1="10", allow_blank=True, showErrorMessage=True,
                           errorTitle="मोबाइल नंबर", error="10 अंक का मोबाइल नंबर डालें, बिना +91 या 0 के।\nEnter the 10-digit mobile number without +91 or 0.",
                           promptTitle="मोबाइल नंबर", prompt="10 अंक, जैसे 9876543210", showInputMessage=True)
    ws.add_data_validation(phone)
    phone.add(f"{letter['phone']}2:{letter['phone']}{last}")
    for r in range(2, last + 1):
        ws[f"{letter['phone']}{r}"].number_format = "@"

    # city: hidden list sheet, dropdown. Any other name can be typed with no error box at all.
    lists = wb.create_sheet("सूची")
    options = city_options(districts)
    for i, opt in enumerate(options, start=1):
        lists.cell(row=i, column=1, value=opt)
    lists.cell(row=1, column=2, value=YES)
    lists.sheet_state = "hidden"
    n = len(options)
    city = DataValidation(type="list", formula1=f"='सूची'!$A$1:$A${n}", allow_blank=True, showErrorMessage=False,
                          promptTitle="शहर", prompt="सूची से चुनें। सूची में न हो तो शहर / कस्बे का नाम खुद लिख दें, या बगल में पिनकोड भरें।", showInputMessage=True)
    ws.add_data_validation(city)
    city.add(f"{letter['city']}2:{letter['city']}{last}")

    # pincode: optional, 6 digits, kept as text
    pin = DataValidation(type="textLength", operator="equal", formula1="6", allow_blank=True, showErrorMessage=True,
                         errorTitle="पिनकोड", error="6 अंक का पिनकोड डालें, जैसे 486001\nEnter the 6-digit pincode, e.g. 486001",
                         promptTitle="पिनकोड", prompt="ज़रूरी नहीं। शहर सूची में न मिले तो 6 अंक का पिनकोड, जैसे 486001", showInputMessage=True)
    ws.add_data_validation(pin)
    pin.add(f"{letter['pincode']}2:{letter['pincode']}{last}")
    for r in range(2, last + 1):
        ws[f"{letter['pincode']}{r}"].number_format = "@"

    # vehicles: only "हाँ" or empty
    yes = DataValidation(type="list", formula1=f'"{YES}"', allow_blank=True, showErrorMessage=True,
                         errorTitle="गाड़ी", error="सिर्फ़ \"हाँ\" चुनें या खाली छोड़ें।\nPick \"हाँ\" or leave empty.")
    ws.add_data_validation(yes)
    a, b = get_column_letter(first_vehicle), get_column_letter(first_vehicle + len(VEHICLE_COLS) - 1)
    yes.add(f"{a}2:{b}{last}")
    for col in range(first_vehicle, first_vehicle + len(VEHICLE_COLS)):
        for r in range(2, last + 1):
            ws.cell(row=r, column=col).alignment = Alignment(horizontal="center")

    if role == "owner":
        count = DataValidation(type="whole", operator="between", formula1="1", formula2="10000", allow_blank=True, showErrorMessage=True,
                               errorTitle="गिनती", error="सिर्फ़ गिनती लिखें, जैसे 12\nNumbers only, e.g. 12")
        ws.add_data_validation(count)
        count.add(f"{letter['count']}2:{letter['count']}{last}")

    # light grid so empty rows look fillable
    for r in range(2, 41):
        for col in range(1, len(cols) + 1):
            ws.cell(row=r, column=col).border = Border(bottom=THIN, right=THIN)

    _help_sheet(wb, role)
    buf = io.BytesIO()
    wb.save(buf)
    return buf.getvalue()


def _help_sheet(wb: Workbook, role: str) -> None:
    hs = wb.create_sheet("कैसे भरें - How to fill", 1)
    hs.column_dimensions["A"].width = 3
    for i in range(2, 13):
        hs.column_dimensions[get_column_letter(i)].width = 15
    lines = [
        ("कैसे भरें", True),
        ("1. पहली शीट \"भरें - Fill here\" में हर व्यक्ति की एक लाइन भरें।", False),
        ("2. मोबाइल नंबर ज़रूरी है: 10 अंक, बिना +91 या 0 के (जैसे 9876543210)। गलत नंबर Excel खुद रोक देगा।", False),
        ("3. नाम हिंदी या English, किसी में भी लिखें।", False),
        ("4. शहर: खाने पर क्लिक करें, तीर से सूची खोलें और चुनें। सूची में न हो तो शहर / कस्बे का नाम खुद लिख दें (हिंदी या English)।", False),
        ("   शहर समझ न आए तो \"पिनकोड\" वाले खाने में 6 अंक का पिनकोड भर दें, जैसे 486001। पिनकोड से जगह अपने आप पता चल जाएगी।", False),
        ("5. गाड़ी: जो गाड़ी चलाते / रखते हैं, उसके नीचे \"हाँ\" चुनें। कई गाड़ियाँ हों तो कई में \"हाँ\"। बाकी खाली छोड़ें।", False),
    ]
    if role == "owner":
        lines.append(("6. कुल गाड़ियाँ: सिर्फ़ गिनती लिखें, जैसे 12।", False))
    lines += [
        ("कोई खाना समझ न आए तो खाली छोड़ दें। सिर्फ़ मोबाइल नंबर से भी काम चल जाएगा।", True),
        ("", False),
        ("How to fill (English)", True),
        ("One row per person on the first sheet. Only the 10-digit mobile number is required.", False),
        ("Name in Hindi or English. City: pick from the dropdown, or type any city / town, or just fill the 6-digit pincode. Vehicles: pick \"हाँ\" under each vehicle, leave others empty.", False),
        ("Not sure about a cell? Leave it empty.", False),
        ("", False),
        ("उदाहरण / Example", True),
    ]
    r = 1
    for text, bold in lines:
        c = hs.cell(row=r, column=2, value=text)
        c.font = Font(bold=bold, size=14 if bold else 12, color="0E5A6B" if bold else "13252A")
        c.alignment = Alignment(wrap_text=True, vertical="center")
        hs.merge_cells(start_row=r, start_column=2, end_row=r, end_column=12)
        hs.row_dimensions[r].height = 24 if text else 10
        r += 1
    # small example table (not imported: it is on this sheet)
    head = ["मोबाइल", "नाम"] + (["फर्म"] if role == "owner" else []) + ["शहर"] + (["कुल गाड़ियाँ"] if role == "owner" else []) + ["ट्रक", "बस", "…"]
    row = ["9876543210", "रमेश कुमार"] + (["कुमार ट्रांसपोर्ट"] if role == "owner" else []) + ["रीवा (Rewa), Madhya Pradesh"] + (["12"] if role == "owner" else []) + [YES, YES, ""]
    for i, h in enumerate(head):
        c = hs.cell(row=r, column=2 + i, value=h)
        c.font = Font(bold=True, color="FFFFFF")
        c.fill = TEAL
    for i, v in enumerate(row):
        c = hs.cell(row=r + 1, column=2 + i, value=v)
        c.fill = YELLOW if i == 0 else SOFT
    city_col = 2 + head.index("शहर")
    hs.column_dimensions[get_column_letter(city_col)].width = 30
