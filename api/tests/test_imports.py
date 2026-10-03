import io

from conftest import make_token

H = {"Authorization": f"Bearer {make_token()}"}
ADMIN = {"id": "11111111-1111-1111-1111-111111111111", "tenant_id": "vahanza", "role": "admin", "blocked": False}


def test_columns_hindi_and_english():
    from vz.importer import map_columns
    assert map_columns(["S.No", "Contact Person", "Mobile No.", "Transport Name", "City", "Vehicle No", "No. of Vehicles", "Vehicle Type"]) == {
        1: "name", 2: "phone", 3: "business_name", 4: "district", 6: "vehicle_count", 7: "vehicles"}
    assert map_columns(["नाम", "मोबाइल नंबर", "शहर", "गाड़ी", "राज्य"]) == {0: "name", 1: "phone", 2: "district", 3: "vehicles", 4: "state"}


def test_phone_cleaning():
    from vz.importer import clean_phone
    assert clean_phone("+91 98765 43210") == ("919876543210", False)
    assert clean_phone(9876543210.0) == ("919876543210", False)
    assert clean_phone("09876543210") == ("919876543210", False)
    assert clean_phone("98765 43210, 9123456789") == ("919876543210", True)
    assert clean_phone("5876543210")[0] is None      # Indian mobiles start with 6-9
    assert clean_phone("9999999999")[0] is None      # all same digit
    assert clean_phone(None) == (None, False)


def test_vehicle_words():
    from vz.importer import clean_vehicles
    assert clean_vehicles("Truck/Hywa, बस और JCB") == (["truck", "bus", "jcb"], False)
    assert clean_vehicles("Tata Ace") == (["pickup"], False)
    assert clean_vehicles("cycle") == ([], True)


def test_xlsx_and_csv_read():
    from openpyxl import Workbook
    from vz.importer import read_table
    wb = Workbook()
    wb.active.append(["नाम", "मोबाइल"])
    wb.active.append(["रमेश", 9876543210])
    b = io.BytesIO()
    wb.save(b)
    assert read_table(b.getvalue())[1] == ["रमेश", 9876543210]
    assert read_table("name;mobile\nA;9876543210\n".encode("utf-8-sig"))[1] == ["A", "9876543210"]
    assert read_table("name\tmobile\nA\t9876543210\n".encode("utf-16"))[1] == ["A", "9876543210"]


def test_preview_statuses(client, db):
    db.respond("from public.profiles where id", ADMIN)
    db.respond("from public.profiles where tenant_id = %s and phone = any", [{"phone": "919811100001"}])
    db.respond("from public.prospects where tenant_id", [{"phone": "919826122222", "joined_at": None}])
    body = "Name,Mobile\nA,9826111111\nB,12345\nC,9826111111\nD,9811100001\nE,9826122222\n"
    r = client.post("/api/v1/admin/imports?role=driver&dry_run=true", headers=H, content=body.encode())
    assert r.status_code == 200
    c = r.json()["counts"]
    assert (c["new"], c["update"], c["on_app"], c["bad"], c["total"]) == (1, 1, 1, 2, 5)
    st = {x["row"]: x["status"] for x in r.json()["rows"]}
    assert st == {2: "new", 3: "bad", 4: "bad", 5: "on_app", 6: "update"}


def test_no_phone_column(client, db):
    db.respond("from public.profiles where id", ADMIN)
    r = client.post("/api/v1/admin/imports?role=driver", headers=H, content=b"name,city\nA,B\n")
    assert r.status_code == 422 and r.json()["detail"]["code"] == "no_phone_column"


def test_non_admin_cannot_import(client, db):
    db.respond("from public.profiles where id", {**ADMIN, "role": "driver"})
    assert client.post("/api/v1/admin/imports?role=driver", headers=H, content=b"mobile\n9876543210\n").status_code == 403


def test_sms_needs_app_url(client, db, monkeypatch):
    from vz.config import get_settings
    monkeypatch.setenv("PUBLIC_APP_URL", "")
    get_settings.cache_clear()
    db.respond("from public.profiles where id", ADMIN)
    r = client.post("/api/v1/admin/prospects/invite-sms", headers=H, json={})
    get_settings.cache_clear()
    assert r.status_code == 409 and r.json()["detail"]["code"] == "no_app_url"


def test_claim_copies_only_for_same_role(db):
    from vz import prospects
    db.respond("update public.prospects set joined_at", {"role": "owner", "name": "S", "business_name": "B", "district": "Rewa", "state": "MP", "vehicles": []})
    assert prospects.claim(db, "vahanza", "u1", "919826122222", "driver") is True
    assert not any("update public.profiles" in c[0] for c in db.calls)
    db.respond("update public.prospects set joined_at", {"role": "driver", "name": "S", "business_name": None, "district": "Rewa", "state": "MP", "vehicles": ["truck", "spaceship"]})
    assert prospects.claim(db, "vahanza", "u2", "919826122222", "driver") is True
    veh = next(c for c in db.calls if "update public.driver_details" in c[0])
    assert veh[1][0] == ["truck"]


def test_invite_text():
    from vz import prospects
    t = prospects.invite_text("driver", "Ramesh Kumar", "Vahanza", "https://x.in/?inv=driver")
    assert t.startswith("नमस्ते Ramesh जी") and t.endswith("https://x.in/?inv=driver")
    assert prospects.whatsapp_url("919876543210", "a b").endswith("919876543210?text=a%20b")


def test_excel_template_round_trip():
    """The downloadable template: dropdown columns come back as clean data."""
    import io
    from openpyxl import load_workbook
    from vz.importer import flag_columns, map_columns, read_table
    from vz.template import build
    wb = load_workbook(io.BytesIO(build("owner")))
    assert wb.sheetnames[0].startswith("भरें") and wb["सूची"].sheet_state == "hidden"
    ws = wb.worksheets[0]
    ws["A2"], ws["B2"], ws["E2"], ws["F2"], ws["G2"], ws["I2"] = "9826199991", "सुरेश", "486001", 12, "हाँ", "हाँ"
    buf = io.BytesIO()
    wb.save(buf)
    rows = read_table(buf.getvalue())
    cols, flags = map_columns(rows[0]), flag_columns(rows[0])
    assert set(cols.values()) == {"phone", "name", "business_name", "district", "pincode", "vehicle_count"}
    assert [flags[i] for i in sorted(flags)] == ["truck", "trailer", "bus", "pickup", "jcb", "tractor", "car", "auto"]
    assert len(rows) == 2


def test_template_lists_all_districts():
    """With the pincode directory loaded, every district is in the city dropdown, not only the curated ones."""
    from vz.template import city_options
    base = city_options()
    more = city_options([("Sidhi", "Madhya Pradesh"), ("Rewari", "Haryana"), ("Rewa", "Madhya Pradesh")])
    assert "Rewari, Haryana" in more and len(more) == len(base) + 1   # Sidhi and Rewa are curated already
    assert not any(o == "Rewa, Madhya Pradesh" for o in more)   # curated "रीवा (Rewa), …" is not repeated


def test_yes_values():
    from vz.importer import is_yes
    assert all(is_yes(v) for v in ["हाँ", "हां", "Yes", "y", 1, "✓", True])
    assert not any(is_yes(v) for v in [None, "", "नहीं", "no", 0])
